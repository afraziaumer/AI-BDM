"""Generates qa/AI-BDM_QA_Test_Plan.xlsx from the test-case data below.
Not part of the application -- a one-off generator script, safe to
re-run any time the test plan needs regenerating after a real code
change. Uses openpyxl only, no formulas (pure tracking data)."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT_NAME, size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="595959")
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PRIORITY_FILL = {
    "Critical": PatternFill("solid", fgColor="F8CBAD"),
    "High": PatternFill("solid", fgColor="FCE4D6"),
    "Medium": PatternFill("solid", fgColor="FFF2CC"),
    "Low": PatternFill("solid", fgColor="E2EFDA"),
}

COLUMNS = [
    ("Test ID", 12),
    ("Module", 20),
    ("Test Case Title", 30),
    ("Preconditions", 26),
    ("Test Steps", 46),
    ("Test Data / Example", 26),
    ("Expected Result", 40),
    ("Priority", 10),
    ("Existing Automated Coverage", 30),
    ("Pass / Fail", 12),
    ("Actual Result / Notes", 26),
]


def write_sheet(wb: Workbook, title: str, subtitle: str, rows: list[tuple]) -> Worksheet:
    ws = wb.create_sheet(title=title[:31])
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:K1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:K2")
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16

    header_row = 4
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 30

    for r, row in enumerate(rows, start=header_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP if c not in (1, 8, 10) else WRAP_CENTER
            cell.border = BORDER
            if c == 8 and value in PRIORITY_FILL:  # Priority column
                cell.fill = PRIORITY_FILL[value]

    last_col = get_column_letter(len(COLUMNS))
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row}"
    return ws


def T(test_id, module, title, pre, steps, data, expected, priority, coverage):
    return (test_id, module, title, pre, steps, data, expected, priority, coverage, "", "")


# ===========================================================================
# READ ME sheet content is built separately (no test-case columns)
# ===========================================================================

def build_readme(wb: Workbook) -> None:
    ws = wb.create_sheet("Read Me")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110

    lines = [
        ("AI-BDM — QA Test Plan", TITLE_FONT),
        ("Full-project test case catalogue handed off for independent QA testing.", SUBTITLE_FONT),
        ("", BODY_FONT),
        ("How this workbook is organized", HEADER_FONT_LOCAL := Font(name=FONT_NAME, bold=True, size=12)),
        ("Each tab after this one is one testable area of the project. Every row is one test "
         "case: what to set up, exact steps, what input to use, and the precise expected "
         "result -- taken from this codebase's actual, verified behavior (its own test suite, "
         "docs/ARCHITECTURE.md, docs/DATA_CONTRACTS.md, docs/backend-handoff/), not invented.", BODY_FONT),
        ("", BODY_FONT),
        ("Columns", HEADER_FONT_LOCAL),
        ("Test ID -- unique, prefixed by area (e.g. STG1-03 = Stage 1, case 3).", BODY_FONT),
        ("Priority -- Critical (blocks core functionality / data integrity), High (a real "
         "documented feature breaks), Medium (degraded behavior, has a fallback), Low "
         "(cosmetic / edge case).", BODY_FONT),
        ("Existing Automated Coverage -- if this exact scenario already has a passing automated "
         "test, its file:test_name is listed. 'None' means this is genuinely new manual "
         "ground -- prioritize those first, since the automated ones are already proven.", BODY_FONT),
        ("Pass/Fail and Actual Result/Notes -- left blank for the tester to fill in.", BODY_FONT),
        ("", BODY_FONT),
        ("Before you start", HEADER_FONT_LOCAL),
        ("1. Read docs/HANDOFF_SUMMARY.md and docs/backend-handoff/KNOWN_LIMITATIONS.md first. "
         "Several rows in the 'Known Limitations' tab of this workbook describe behavior that "
         "is EXPECTED, not a bug -- filing a defect against those wastes everyone's time.", BODY_FONT),
        ("2. Environment: Python 3.12, Windows (developed/tested on) or Linux/macOS "
         "(not separately verified). See README.md for exact install commands.", BODY_FONT),
        ("3. Run the automated suite yourself first: `pytest -q` (offline, no API keys, ~115 "
         "tests) and, only when you intend to spend real provider quota, "
         "`pytest -m live tests/integration` (4 tests, costs 1 real Serper search credit).", BODY_FONT),
        ("4. Some test cases require real API keys (Groq, Serper, ScrapingBee, Apify) in a "
         "local .env -- see .env.example. Never commit a real .env file.", BODY_FONT),
        ("5. Two ways to exercise the pipeline: `python main.py --query \"...\"` (full 8-stage "
         "CLI) or the REST API (`python -m uvicorn api:app --reload --port 8000`, partial: "
         "Stages 1-2 only, plus the async job endpoints under /api/v1/jobs/*).", BODY_FONT),
        ("", BODY_FONT),
        ("Tabs in this workbook", HEADER_FONT_LOCAL),
        ("Setup & Environment -- installing, configuring, and proving a clean environment works.", BODY_FONT),
        ("Pipeline Stages 1-8 -- the core CLI pipeline, stage by stage.", BODY_FONT),
        ("Sync API -- the blocking REST endpoints (/health, /pipeline/run, /leads).", BODY_FONT),
        ("Async Job API -- the queued job endpoints (/jobs, progress, cancel, resume, Idempotency-Key).", BODY_FONT),
        ("Data & Storage Contracts -- what gets written to disk and its guarantees.", BODY_FONT),
        ("Error Handling & Resilience -- provider failures, retries, circuit breakers.", BODY_FONT),
        ("Security -- auth, secrets, fail-closed behavior.", BODY_FONT),
        ("Known Limitations -- documented, accepted gaps. Confirm the documented behavior, don't file bugs.", BODY_FONT),
    ]
    r = 1
    for text, font in lines:
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30 if len(text) > 90 else (24 if text else 8)
        r += 1


# ===========================================================================
# SETUP & ENVIRONMENT
# ===========================================================================

SETUP_ROWS = [
    T("SETUP-01", "Setup", "Fresh virtual environment creation",
      "Python 3.12 installed", "1. `py -3.12 -m venv .venv`\n2. Activate it\n3. `python -m pip install --upgrade pip`",
      "N/A", "venv created and activated with no errors", "Critical", "None (manual)"),
    T("SETUP-02", "Setup", "Install pinned dependencies",
      "Fresh venv active", "`pip install -r requirements.txt`",
      "requirements.txt (repo root)", "All packages install cleanly; `wappalyzer` is NOT installed as a pip package (it's vendored)", "Critical", "docs/backend-handoff/TEST_RESULTS.txt (fresh-install verification)"),
    T("SETUP-03", "Setup", "Post-install segfault-avoidance step",
      "SETUP-02 done", "`pip uninstall -y pyarrow datasets`",
      "N/A", "Both packages removed; safe because nothing in this codebase imports either directly", "High", "None (manual)"),
    T("SETUP-04", "Setup", "Vendored Wappalyzer resolves correctly",
      "SETUP-02/03 done", "1. `python -c \"import wappalyzer; print(wappalyzer.__file__)\"`",
      "N/A", "Path resolves to vendor/wappalyzer/__init__.py, NOT site-packages -- proves the sys.path shim in tech_stack.py works", "Critical", "tests/unit/test_wappalyzer_patches.py"),
    T("SETUP-05", "Setup", ".env configuration from template",
      "SETUP-02 done", "1. Copy .env.example to .env\n2. Fill in real Groq + Serper keys (minimum required)\n3. `python main.py --query \"find 1 dental clinic in Lahore\"`",
      ".env.example", "Pipeline runs to completion; no 'missing API key' errors for the two required services", "Critical", "None (manual, needs real keys)"),
    T("SETUP-06", "Setup", "Missing required GROQ key fails clearly",
      "GROQ key removed/blank from .env", "Run `python main.py --query \"find 1 dental clinic\"`",
      "N/A", "Fails with a clear intent_failed error, not a silent hang or generic crash", "High", "None (manual)"),
    T("SETUP-07", "Setup", "Missing required SERPER key fails clearly",
      "SERPER key removed/blank, GROQ present", "Run `python main.py --query \"find 1 dental clinic\"`",
      "N/A", "Planning succeeds (LLM works), discovery fails clearly -- no businesses found, error is legible not a raw traceback", "High", "None (manual)"),
    T("SETUP-08", "Setup", "Offline smoke suite passes with zero network access",
      "SETUP-02 done, no .env needed", "1. Disconnect network (or just don't set any API keys)\n2. `pytest -q tests/smoke`",
      "N/A", "14/14 pass, no network calls attempted", "Critical", "tests/smoke/ (self-verifying)"),
    T("SETUP-09", "Setup", "Full offline + mocked suite passes",
      "SETUP-02 done", "`pytest -q`",
      "N/A", "115 passed, 4 deselected (the live-provider tests), zero failures", "Critical", "Entire tests/ tree except tests/integration"),
    T("SETUP-10", "Setup", "Live provider credential suite (costs real quota)",
      "Real API keys in .env for Groq, ScrapingBee, Apify, Serper", "`pytest -m live tests/integration`",
      "N/A", "4 passed -- confirms all 4 provider credentials are valid. Serper check spends 1 real search credit by design.", "Medium", "tests/integration/test_live_providers.py"),
    T("SETUP-11", "Setup", "Windows console UTF-8 handling",
      "Running on Windows", "Run a query where the LLM's output is likely to include an em-dash or smart quote (most natural-language answers)",
      "Any query hitting Stage 4 (find_and_filter intent)", "No UnicodeEncodeError -- main.py reconfigures stdout/stderr to UTF-8 at startup", "Medium", "None (manual, Windows-specific)"),
]

# ===========================================================================
# PIPELINE STAGES 1-8
# ===========================================================================

STAGE_ROWS = [
    # --- Stage 1: Query Planning ---
    T("STG1-01", "Stage 1: Query Planning", "Valid natural-language query produces a full plan",
      "GROQ key configured", "`python main.py --query \"find 3 dental clinics in Islamabad with no online booking\"`",
      "See example above", "Plan JSON has all fields populated: geo_location, broad_industry, search_query, result_limit, search_type, intent, etc.", "Critical", "tests/integration (live only)"),
    T("STG1-02", "Stage 1: Query Planning", "Vague query triggers clarification gate",
      "GROQ key configured", "`python main.py --query \"find some businesses\"`",
      "\"find some businesses\"", "Pipeline stops BEFORE Stage 2; needs_clarification=true; numbered questions printed; NO Serper/scrape calls made", "Critical", "None (manual, needs live LLM judgment)"),
    T("STG1-03", "Stage 1: Query Planning", "Clarification answer causes re-plan and continuation",
      "STG1-02 reproduced", "Type an answer at the \"Your answer:\" prompt",
      "e.g. \"Karachi, restaurants\"", "Pipeline re-plans with the added detail and proceeds to Stage 2 automatically, in the same session", "High", "None (manual)"),
    T("STG1-04", "Stage 1: Query Planning", "3 unanswered clarification rounds proceeds with best guess",
      "STG1-02 reproduced", "Press Enter (empty answer) 3 times in a row",
      "N/A", "After round 3, pipeline proceeds anyway using its best-guess plan -- never blocks forever", "Medium", "None (manual)"),
    T("STG1-05", "Stage 1: Query Planning", "Explicit count in query is honored exactly",
      "GROQ key configured", "`python main.py --query \"give me 7 marinas in Miami\"`",
      "\"give me 7 marinas in Miami\"", "plan.result_limit=7, count_explicit=true; pipeline works toward exactly 7 qualified leads, not the default", "High", "tests/unit/test_query_input.py (resolve_effective_limit logic)"),
    T("STG1-06", "Stage 1: Query Planning", "No count in query uses single-pass default",
      "GROQ key configured", "`python main.py --query \"find marinas in Miami\"`",
      "\"find marinas in Miami\"", "count_explicit=false; discovery_mode=single_pass; only 1 discovery round runs, no query-variation expansion", "Medium", "tests/unit/test_query_input.py"),
    T("STG1-07", "Stage 1: Query Planning", "Named single business triggers specific search",
      "GROQ key configured", "`python main.py --query \"find nike.com\"`",
      "\"find nike.com\"", "search_type=specific, target_domain=nike.com; only that one site is targeted", "High", "None (manual)"),
    T("STG1-08", "Stage 1: Query Planning", "Tech-related keyword sets needs_tech_stack",
      "GROQ key configured", "`python main.py --query \"find clinics using WordPress in Lahore\"`",
      "\"...using WordPress...\"", "plan.needs_tech_stack=true -> Stage 5 runs for each committed business", "Medium", "None (manual)"),
    T("STG1-09", "Stage 1: Query Planning", "Groq total failure stops pipeline cleanly",
      "Simulate: invalid GROQ key, or provider outage", "Run any query", "N/A",
      "summary['error'] starts with 'intent_failed:'; pipeline stops; NO partial/fabricated plan is used", "Critical", "None (manual, needs simulated failure)"),
    T("STG1-10", "Stage 1: Query Planning", "Policy-violating query is blocked before any paid call",
      "GROQ key configured", "Submit an obviously policy-violating query (e.g. asks for something clearly against usage policy)",
      "N/A", "summary['blocked']=true, block_reason set, error='blocked_by_moderation'; zero Serper/ScrapingBee calls made", "Critical", "None (manual, needs live moderation judgment)"),
    T("STG1-11", "Stage 1: Query Planning", "Benign query is never falsely blocked",
      "GROQ key configured", "Submit 5-10 ordinary business-research queries",
      "Various normal queries", "None of them are blocked -- confirms moderation isn't over-triggering on legitimate research", "High", "None (manual)"),

    # --- Stage 2: Discovery & Scraping ---
    T("STG2-01", "Stage 2: Discovery & Scraping", "General search discovers and commits qualified businesses",
      "GROQ + SERPER keys configured", "`python main.py --query \"find 3 dental clinics in Islamabad\"`",
      "N/A", "storage/<domain>/ created per business; crawl_index.csv gets rows; summary['qualified'] has ~3 entries", "Critical", "None (manual, live)"),
    T("STG2-02", "Stage 2: Discovery & Scraping", "Specific search returns cached record without re-scraping",
      "A domain already committed in storage/ from a prior run", "`python main.py --query \"find <that-domain>\"`",
      "An already-crawled domain", "Log shows 'already in store -- returning existing'; no new HTTP fetch to that domain; existing scoring_version/keywords re-validated", "High", "None (manual)"),
    T("STG2-03", "Stage 2: Discovery & Scraping", "New specific-domain search scrapes exactly once",
      "Domain NOT previously in storage/", "`python main.py --query \"find <new-domain>\"`",
      "A domain never crawled before", "One crawl job created, storage/<domain>/ committed", "High", "None (manual)"),
    T("STG2-04", "Stage 2: Discovery & Scraping", "Exclude keyword found on page still commits the business",
      "Query with an exclude keyword the target site actually has", "Query for an industry with an explicit exclude keyword (e.g. \"...with no CRM\") against a site that DOES have a CRM",
      "N/A", "qualification_status=excluded_by_keyword; business is STILL committed to storage (never silently dropped), excluded_keywords column populated", "High", "None (manual)"),
    T("STG2-05", "Stage 2: Discovery & Scraping", "ScrapingBee premium tier escalates on native-fetch block",
      "ScrapingBee key configured; target site known to block plain requests (Cloudflare/captcha)", "Crawl a site known to block native fetch",
      "N/A", "Native fetch fails, ScrapingBee premium tier is invoked automatically, page is successfully retrieved", "High", "None (manual, needs a real blocking site)"),
    T("STG2-06", "Stage 2: Discovery & Scraping", "Circuit breaker opens after 2 consecutive premium failures",
      "ScrapingBee key configured", "Crawl a domain whose premium fetch fails twice in a row (e.g. simulate by pointing at an unreachable host)",
      "N/A", "Log: \"Premium scraper failed 2 consecutive time(s)... opening circuit breaker\"; remaining pages for that domain use native-fetch-only, no further premium attempts", "High", "None (manual)"),
    T("STG2-07", "Stage 2: Discovery & Scraping", "Serper discovery failure is retried, never cached as empty",
      "Simulate network interruption during discovery", "Interrupt network mid-discovery-search, then re-run the same query",
      "N/A", "First run: no false 'zero results' cached. Second run: discovery is attempted again fresh", "Critical", "tests/mocked/test_cache_correctness.py (equivalent logic for Maps/reviews)"),
    T("STG2-08", "Stage 2: Discovery & Scraping", "Directory site mines and queues real businesses",
      "A query whose top result is a directory/association listing page", "Query an industry where the first search result is a directory site",
      "N/A", "That result is classified 'directory'; its discovered_businesses are queued and each crawled individually in later rounds", "Medium", "None (manual)"),
    T("STG2-09", "Stage 2: Discovery & Scraping", "Query-variation expansion fires only for explicit-count queries once exhausted",
      "An explicit-count query for a niche industry/location combo with few real matches", "`python main.py --query \"give me 20 <very niche business type> in <small town>\"`",
      "N/A", "Once the original query is exhausted, log shows LLM-generated query variations being tried; shortfall reported honestly if still not met (never padded)", "Medium", "None (manual)"),
    T("STG2-10", "Stage 2: Discovery & Scraping", "Single-pass mode never expands beyond 1 round",
      "A query with no explicit count", "`python main.py --query \"find dental clinics in Lahore\"` (no number)",
      "N/A", "Exactly 1 discovery round runs regardless of how few/many qualified leads it finds; no query-variation expansion triggered", "Medium", "None (manual)"),
    T("STG2-11", "Stage 2: Discovery & Scraping", "Re-run of same query skips already-committed businesses",
      "A query already run once to completion", "Run the exact same query a second time",
      "N/A", "Log shows cached-relevance reuse for already-committed domains under the same scoring_version/industry/geo; no re-scrape", "High", "docs/DATA_CONTRACTS.md re-run/idempotency section"),
    T("STG2-12", "Stage 2: Discovery & Scraping", "Same domain re-validated against a different query",
      "A domain committed under Query A (e.g. 'marinas in Miami')", "Run Query B against the same domain with a different industry/geo (e.g. 'law firms in Chicago')",
      "N/A", "Domain is re-validated against the NEW industry/geo, not blindly assumed qualified from the earlier run", "Medium", "None (manual)"),
    T("STG2-13", "Stage 2: Discovery & Scraping", "robots.txt Disallow rules are NOT enforced (documented gap)",
      "A target site with a restrictive robots.txt Disallow rule on a page this pipeline would otherwise crawl", "Crawl that site",
      "N/A", "EXPECTED (not a bug): the page IS fetched despite the Disallow rule -- robots.txt is only parsed for Sitemap: discovery, never enforced. See Known Limitations tab.", "Low", "None (manual) -- see docs/backend-handoff/KNOWN_LIMITATIONS.md"),
    T("STG2-14", "Stage 2: Discovery & Scraping", "Job cancellation stops before next round, keeps committed work",
      "Submit an async job via POST /api/v1/jobs for a multi-round query", "Call POST /api/v1/jobs/{job_id}/cancel partway through",
      "N/A", "Job stops starting new discovery rounds; status becomes 'cancelled'; businesses already committed before cancellation remain in storage/", "High", "tests/mocked/test_job_lifecycle.py::test_job_can_be_cancelled_mid_run"),

    # --- Stage 3: Route/Page Selection ---
    T("STG3-01", "Stage 3: Route Selection", "Multi-page domain uses search-first retrieval",
      "A committed domain with 3+ crawled pages, find_and_filter query", "Run a query with intent=find_and_filter against a multi-page business",
      "N/A", "Log shows BM25+embedding search-first selection (e.g. \"1 page(s) -> 1 candidate(s)\" style), no LLM call for routing in the common case", "Medium", "None (manual)"),
    T("STG3-02", "Stage 3: Route Selection", "Single-page domain uses heuristic selection",
      "A committed domain with exactly 1 crawled page", "Run a find_and_filter query against it",
      "N/A", "URL/title heuristic selection used, no full search index built for one document", "Low", "None (manual)"),
    T("STG3-03", "Stage 3: Route Selection", "LLM fallback fires when retrieval finds nothing usable",
      "A domain whose pages genuinely don't match the query's intent", "Run a find_and_filter query where local retrieval would return nothing",
      "N/A", "TaskType.ROUTE_PLANNING LLM call fires as fallback (rare path) -- confirm via logs", "Low", "None (manual)"),

    # --- Stage 4: Final Reasoning ---
    T("STG4-01", "Stage 4: Final Reasoning", "find_and_filter query triggers Stage 4 with cited answer",
      "GROQ key configured", "`python main.py --query \"find 3 dental clinics in Islamabad with no online booking\"`",
      "N/A", "answer_query() returns {answer, source_pages, confidence}; source_pages cites real filenames the answer was grounded in", "Critical", "None (manual, live)"),
    T("STG4-02", "Stage 4: Final Reasoning", "Plain find query skips Stage 4 entirely",
      "GROQ key configured", "`python main.py --query \"find 3 dental clinics in Islamabad\"` (no filter clause)",
      "N/A", "intent=find (not find_and_filter); Stage 4 never invoked for these businesses", "Medium", "None (manual)"),
    T("STG4-03", "Stage 4: Final Reasoning", "No retrievable pages degrades gracefully",
      "A committed business with zero usable crawled text", "Run a find_and_filter query against it",
      "N/A", "confidence='low', reason='no_retrievable_pages' -- never raises an exception to the caller", "High", "None (manual)"),
    T("STG4-04", "Stage 4: Final Reasoning", "Insufficient evidence produces honest 'not enough evidence' answer",
      "A business whose pages don't address the specific filter asked about", "Ask a find_and_filter question the site's content can't actually answer",
      "N/A", "Answer explicitly states there isn't enough evidence, rather than fabricating a claim", "Critical", "None (manual, live)"),

    # --- Stage 5: Tech-Stack Detection ---
    T("STG5-01", "Stage 5: Tech-Stack Detection", "needs_tech_stack=true produces a full profile",
      "GROQ key configured, vendored wappalyzer installed", "`python main.py --query \"find businesses using WordPress in Lahore\"`",
      "N/A", "storage/<domain>/tech_stack.json written with raw_wappalyzer, normalized_tech_stack, sales_signals all populated", "High", "None (manual, live)"),
    T("STG5-02", "Stage 5: Tech-Stack Detection", "Stage skipped when not requested",
      "A plain query with no tech-stack intent", "`python main.py --query \"find 3 dental clinics\"`",
      "N/A", "No tech_stack.json written for those businesses (or it's empty/absent, per needs_tech_stack=false)", "Medium", "None (manual)"),
    T("STG5-03", "Stage 5: Tech-Stack Detection", "Partial engine failure still yields a profile",
      "Any tech-stack-requiring query", "Simulate one engine failing (e.g. block DNS lookups) during a tech-stack scan",
      "N/A", "The other 3 engines (legacy Wappalyzer, extended cookie/DOM/JS, robots.txt) still populate a partial profile -- one engine's failure never breaks the whole result", "Medium", "None (manual)"),
    T("STG5-04", "Stage 5: Tech-Stack Detection", "Known CMS is correctly fingerprinted",
      "A real WordPress or Shopify site as the query target", "Run tech-stack detection against a known WordPress site",
      "A real WordPress URL", "WordPress correctly appears in the detected technologies", "High", "None (manual, live)"),
    T("STG5-05", "Stage 5: Tech-Stack Detection", "Symfony fingerprint regex works (previously-bugged pattern)",
      "A real Symfony-based site", "Run tech-stack detection against a known Symfony site",
      "A real Symfony URL", "Symfony is correctly detected -- confirms the vendored copy's regex patch is applied", "Medium", "tests/unit/test_wappalyzer_patches.py (compile-level, not live-detection)"),
    T("STG5-06", "Stage 5: Tech-Stack Detection", "60-day-old profile is refreshed on next read",
      "A tech_stack.json with last_scanned older than 60 days (edit the file's timestamp manually)", "Query the same business again with needs_tech_stack=true",
      "Manually backdated tech_stack.json", "Profile is re-scanned rather than reused, since it's past TECH_PROFILE_TTL_DAYS", "Low", "None (manual)"),

    # --- Stage 6: RAG / Evidence ---
    T("STG6-01", "Stage 6: RAG/Evidence", "--no-rag flag skips evidence indexing",
      "GROQ + SERPER keys configured", "`python main.py --query \"find 3 dental clinics in Islamabad\" --no-rag`",
      "N/A", "rag/.chroma/ is not updated for this run; Stage 4 (if triggered) still runs on raw page text, just without the vector index", "Medium", "None (manual)"),
    T("STG6-02", "Stage 6: RAG/Evidence", "Website and review content embedded into separate ranked categories",
      "A business with both scraped pages AND harvested reviews", "Run a query that reaches Stage 6 for such a business",
      "N/A", "Two independently-ranked chunk categories exist -- a large site's chunks don't drown out the business's review signal", "Medium", "None (manual)"),
    T("STG6-03", "Stage 6: RAG/Evidence", "Negation-aware scoring distinguishes has/doesn't-have/planned",
      "A business whose page text says something like \"we don't have a CRM yet, planning one for next year\"", "Run a find_and_filter query asking about CRM presence against that business",
      "N/A", "The answer correctly reflects the nuance (doesn't have one now, one is planned) rather than a naive keyword match saying \"has a CRM\"", "High", "None (manual, live)"),
    T("STG6-04", "Stage 6: RAG/Evidence", "Deleting rag/.chroma/ triggers a clean rebuild",
      "An existing rag/.chroma/ directory", "1. Delete rag/.chroma/\n2. Run any query that reaches Stage 6",
      "N/A", "Index rebuilds from already-scraped page/review text with no re-scraping required", "Low", "None (manual)"),

    # --- Stage 7: Google Maps Enrichment ---
    T("STG7-01", "Stage 7: Maps Enrichment", "--no-maps flag skips Maps enrichment",
      "GROQ + SERPER keys configured", "`python main.py --query \"find 3 dental clinics in Islamabad\" --no-maps`",
      "N/A", "leads_with_maps.csv is not produced/updated for this run; leads_clean.csv still has full data", "Medium", "None (manual)"),
    T("STG7-02", "Stage 7: Maps Enrichment", "Address match confirms a listing",
      "A business with a distinctive, matchable street address", "Run a query for a business with a clear physical address",
      "N/A", "leads_with_maps.csv row has maps_matched=true, match_basis='address'", "High", "tests/mocked/test_maps_matching.py"),
    T("STG7-03", "Stage 7: Maps Enrichment", "Phone match confirms when address is unavailable/ambiguous",
      "A business with a matchable phone but weak/shared address", "Run a query for such a business",
      "N/A", "match_basis='phone', tolerant of formatting differences (spaces/dashes/parens)", "High", "tests/mocked/test_maps_matching.py::test_phone_match_tolerant_of_formatting"),
    T("STG7-04", "Stage 7: Maps Enrichment", "Name-only fallback only fires for a single unambiguous candidate",
      "A business where Serper Places returns exactly 1 result, no address/phone match possible", "Run such a query",
      "N/A", "match_basis='name_only', confidence='low' -- and this ONLY happens when there's exactly one candidate, never when there are multiple", "Medium", "tests/mocked/test_maps_matching.py"),
    T("STG7-05", "Stage 7: Maps Enrichment", "Serper Places failure is retried, not cached as no-match",
      "Simulate a network interruption during Stage 7", "Interrupt network mid-Maps-lookup, then re-run",
      "N/A", "First run: business not permanently marked unmatched. Second run: Maps lookup is retried fresh", "Critical", "tests/mocked/test_cache_correctness.py::test_google_maps_transient_failure_is_not_cached"),
    T("STG7-06", "Stage 7: Maps Enrichment", "Genuine zero-results is cached correctly",
      "A business genuinely absent from Google Maps", "Run a query for a business with no real Maps presence",
      "N/A", "maps_matched=false, cached as a real answer (not retried needlessly next time)", "Medium", "tests/mocked/test_cache_correctness.py::test_google_maps_genuine_empty_response_is_cached"),

    # --- Stage 8: Accuracy Audit ---
    T("STG8-01", "Stage 8: Accuracy Audit", "Runs automatically at end of main.py",
      "Any completed run", "`python main.py --query \"find 3 dental clinics in Islamabad\"`",
      "N/A", "accuracy_report.txt is written/overwritten as the final step, with no separate flag needed to enable it", "Critical", "tests/smoke/test_offline_smoke.py"),
    T("STG8-02", "Stage 8: Accuracy Audit", "Standalone re-check works without re-scraping",
      "leads_with_maps.csv already exists from a prior run", "`python accuracy_check.py --geo \"Islamabad\"`",
      "N/A", "Report regenerated from already-committed data; no new scraping/API calls made", "Medium", "None (manual)"),
    T("STG8-03", "Stage 8: Accuracy Audit", "Malformed field is flagged",
      "A fixture/business with an invalid phone or out-of-range rating", "Run accuracy check against such data",
      "tests/fixtures/leads_with_maps_sample.csv", "[field] flag appears in the report for the malformed value", "High", "tests/smoke/test_offline_smoke.py::test_field_validity_passes_clean_fixture (and its failure-case sibling)"),
    T("STG8-04", "Stage 8: Accuracy Audit", "Cross-source identity mismatch is flagged",
      "A business whose name differs materially between its own site and its Maps/review listing", "Run accuracy check against such a business",
      "N/A", "[identity] flag appears, catching a likely wrong-business match", "High", "None (manual)"),
    T("STG8-05", "Stage 8: Accuracy Audit", "Reddit mention not referencing the business is flagged",
      "A Reddit review record whose text never mentions the business name or geo", "Run accuracy check against such a fixture",
      "N/A", "[review] flag appears for that Reddit mention specifically", "Medium", "None (manual)"),
    T("STG8-06", "Stage 8: Accuracy Audit", "Google Reviews are deliberately NOT checked for name-mention",
      "A Google Reviews record with review text that never says the business's name", "Run accuracy check against it",
      "N/A", "EXPECTED: no [review] flag for Google Reviews on this basis -- listing-based platforms are deliberately excluded from this specific check (see docs/ARCHITECTURE.md Stage 8)", "Low", "None (manual) -- confirms intended design, not a bug"),
    T("STG8-07", "Stage 8: Accuracy Audit", "Non-English review text is skipped, not false-flagged",
      "A review in a language other than English/Roman-Urdu (e.g. genuine Urdu script, or French)", "Run accuracy check against such a review",
      "N/A", "Correctly skipped by language detection, no false [review] flag", "Medium", "None (manual)"),
    T("STG8-08", "Stage 8: Accuracy Audit", "Roman Urdu review text is misdetected (known gap)",
      "A Roman-Urdu-written review, e.g. \"Sindh Balochistan ka idk k Kya criteria hota...\"", "Run accuracy check against such a review",
      "Roman Urdu text sample", "EXPECTED (not a bug): langdetect reports this as English, so the relevance check may false-flag it. See Known Limitations tab.", "Low", "None (manual) -- confirms documented gap"),
]

# ===========================================================================
# SYNC API
# ===========================================================================

SYNC_API_ROWS = [
    T("API-01", "Sync API", "Health endpoint is public (no auth required)",
      "API running: `python -m uvicorn api:app --port 8000`", "`curl http://127.0.0.1:8000/api/v1/health`",
      "N/A", "200 OK, {status: ok, serper_key, premium_scraper_key} -- no X-API-Key header needed", "Critical", "None (manual)"),
    T("API-02", "Sync API", "Health endpoint accurately reports configured keys",
      "API running with only SERPER key set, not ScrapingBee", "`curl http://127.0.0.1:8000/api/v1/health`",
      "N/A", "serper_key='set', premium_scraper_key='missing'", "Medium", "None (manual)"),
    T("API-03", "Sync API", "Protected endpoint rejects missing API key",
      "API running, AIBDM_API_KEY set on server", "`curl -X POST http://127.0.0.1:8000/api/v1/pipeline/run -d '{\"query\":\"find 3 clinics\"}'` (no header)",
      "N/A", "401 Unauthorized", "Critical", "None (manual)"),
    T("API-04", "Sync API", "Protected endpoint rejects wrong API key",
      "API running", "Same as above with `-H \"X-API-Key: wrong-key\"`",
      "N/A", "401 Unauthorized", "Critical", "None (manual)"),
    T("API-05", "Sync API", "Server without AIBDM_API_KEY fails closed",
      "API started with AIBDM_API_KEY unset", "`curl -X POST .../pipeline/run -H \"X-API-Key: anything\" ...`",
      "N/A", "503 Service Unavailable -- protected endpoints refuse ALL calls when unconfigured, never run open", "Critical", "None (manual)"),
    T("API-06", "Sync API", "Valid request runs the pipeline and returns the summary",
      "Correct API key, GROQ+SERPER configured", "`curl -X POST .../pipeline/run -H \"X-API-Key: $KEY\" -d '{\"query\":\"find 3 dental clinics in Islamabad\",\"concurrency\":5}'`",
      "examples/sample_request.json", "200 OK, body matches examples/sample_response.json's shape", "Critical", "None (manual, live)"),
    T("API-07", "Sync API", "Query too short is rejected with 422",
      "Correct API key", "POST with `{\"query\": \"ab\"}` (< 3 chars)",
      "N/A", "422 Unprocessable Entity, Pydantic validation error", "Medium", "tests/unit/test_query_input.py::test_pipeline_request_rejects_too_short_query"),
    T("API-08", "Sync API", "Concurrency out of range is rejected with 422",
      "Correct API key", "POST with `{\"query\": \"find clinics\", \"concurrency\": 50}` (> 20)",
      "N/A", "422 Unprocessable Entity", "Medium", "tests/unit/test_query_input.py::test_pipeline_request_rejects_concurrency_out_of_range"),
    T("API-09", "Sync API", "Moderation-blocked query returns 400",
      "Correct API key", "POST a policy-violating query",
      "N/A", "400 Bad Request, generic policy-violation message (no internal category leaked)", "High", "None (manual, live)"),
    T("API-10", "Sync API", "Intent/LLM failure returns 502",
      "Correct API key, simulate Groq outage", "POST any valid-shaped query during a Groq outage",
      "N/A", "502 Bad Gateway", "High", "None (manual)"),
    T("API-11", "Sync API", "Idempotency-Key replays identical success outcome",
      "Correct API key", "1. POST /pipeline/run with header `Idempotency-Key: test-1`\n2. POST the exact same request again with the same header",
      "Same query + same Idempotency-Key", "Second response body is byte-identical to the first; pipeline does NOT run a second time", "High", "tests/mocked/test_api_jobs.py::test_pipeline_run_idempotency_key_replays_without_rerunning"),
    T("API-12", "Sync API", "Idempotency-Key replays identical error outcome",
      "Correct API key, a query that will error (e.g. simulated LLM failure)", "POST twice with the same Idempotency-Key",
      "N/A", "Both calls return the SAME status code and body (e.g. both 502 with identical detail)", "Medium", "tests/mocked/test_api_jobs.py::test_pipeline_run_idempotency_key_replays_error_outcome_too"),
    T("API-13", "Sync API", "No Idempotency-Key means every call re-runs",
      "Correct API key", "POST the same query twice with NO Idempotency-Key header",
      "N/A", "Pipeline runs twice independently (two separate full executions)", "Medium", "tests/mocked/test_api_jobs.py"),
    T("API-14", "Sync API", "Leads list, first page, default size",
      "Correct API key, some leads already in crawl_index.csv", "`curl .../api/v1/leads -H \"X-API-Key: $KEY\"`",
      "N/A", "200 OK, {items: [...], next_cursor}. Default page size 50.", "High", "None (manual)"),
    T("API-15", "Sync API", "Leads list respects limit parameter",
      "Correct API key, 5+ leads present", "`curl \".../api/v1/leads?limit=2\" -H \"X-API-Key: $KEY\"`",
      "N/A", "Exactly 2 items returned, next_cursor is non-null if more than 2 exist", "High", "None (manual)"),
    T("API-16", "Sync API", "Cursor correctly advances to the next page",
      "Correct API key, 5+ leads present", "1. GET with limit=2, note next_cursor\n2. GET again with `cursor=<that value>&limit=2`",
      "N/A", "Second call returns items 3-4, not a repeat of items 1-2", "Critical", "None (manual)"),
    T("API-17", "Sync API", "Last page returns a null next_cursor",
      "Correct API key", "Page through all leads until the final page",
      "N/A", "Final page's next_cursor is null, signaling no more data", "High", "None (manual)"),
    T("API-18", "Sync API", "Invalid/tampered cursor is rejected",
      "Correct API key", "`curl \".../api/v1/leads?cursor=not-a-real-cursor\" -H \"X-API-Key: $KEY\"`",
      "N/A", "400 Bad Request -- fails closed, doesn't silently reset to page 1", "Medium", "None (manual)"),
    T("API-19", "Sync API", "Leads count is accurate",
      "Correct API key, known number of leads present", "`curl .../api/v1/leads/count -H \"X-API-Key: $KEY\"`",
      "N/A", "{count: N} matches the actual number of rows in crawl_index.csv", "Medium", "None (manual)"),
]

# ===========================================================================
# ASYNC JOB API
# ===========================================================================

JOB_API_ROWS = [
    T("JOB-01", "Async Job API", "Submitting a valid job returns 202 queued",
      "Correct API key", "POST /api/v1/jobs with a valid JobRequest body (job_id, tenant_ref, correlation_id, requested_at, target)",
      "See docs/backend-handoff/schemas/request.schema.json", "202 Accepted, status='queued'", "Critical", "tests/mocked/test_api_jobs.py::test_submit_status_result_lifecycle"),
    T("JOB-02", "Async Job API", "Duplicate job_id is rejected",
      "A job already submitted with job_id X", "POST /api/v1/jobs again with the same job_id X",
      "N/A", "409 Conflict", "Critical", "tests/mocked/test_api_jobs.py::test_duplicate_submit_returns_409"),
    T("JOB-03", "Async Job API", "Status endpoint reflects live progress while running",
      "A job in progress", "GET /api/v1/jobs/{job_id} repeatedly while it runs",
      "N/A", "status='running', stage/completed_units/total_units/message_code update over time, last_heartbeat_at advances", "High", "tests/mocked/test_job_lifecycle.py::test_job_reports_progress_before_completing"),
    T("JOB-04", "Async Job API", "Unknown job_id returns 404 on every job route",
      "N/A", "GET /api/v1/jobs/nope, GET /api/v1/jobs/nope/result, POST /api/v1/jobs/nope/cancel, POST /api/v1/jobs/nope/resume",
      "job_id='nope'", "All four return 404", "Critical", "tests/mocked/test_api_jobs.py::test_unknown_job_id_returns_404, test_resume_unknown_original_returns_404"),
    T("JOB-05", "Async Job API", "Result endpoint returns 202 while still running",
      "A job still in progress", "GET /api/v1/jobs/{job_id}/result",
      "N/A", "202 Accepted, {status, detail: 'Job is still running; poll again.'} -- not an error", "High", "tests/mocked/test_api_jobs.py::test_submit_status_result_lifecycle"),
    T("JOB-06", "Async Job API", "Result endpoint returns the full JobResult once terminal",
      "A completed job", "GET /api/v1/jobs/{job_id}/result",
      "N/A", "200 OK, full JobResult body validating against result.schema.json", "Critical", "tests/mocked/test_api_jobs.py + tests/contract/test_job_schemas.py"),
    T("JOB-07", "Async Job API", "Cancel on a running job stops it and preserves partial results",
      "A job with multiple rounds in progress", "POST /api/v1/jobs/{job_id}/cancel while it's between rounds",
      "N/A", "{cancelled: true}; eventual status='cancelled'; any businesses already committed remain in storage/", "Critical", "tests/mocked/test_api_jobs.py::test_cancel_endpoint"),
    T("JOB-08", "Async Job API", "Cancel on an already-terminal job is not an error",
      "A job that already completed", "POST /api/v1/jobs/{job_id}/cancel on it",
      "N/A", "200 OK, {cancelled: false} -- not a 4xx/5xx", "Medium", "tests/mocked/test_job_lifecycle.py"),
    T("JOB-09", "Async Job API", "company_size in target produces an explicit warning, not silence",
      "A job request with target.company_size set", "Submit such a job, wait for result",
      "target.company_size={min_employees:10,max_employees:250}", "JobResult.warnings includes a note that company_size isn't supported -- never silently dropped", "High", "tests/unit/test_job_translation.py"),
    T("JOB-10", "Async Job API", "features.maps=true produces a scope warning",
      "A job request with features.maps=true", "Submit such a job",
      "N/A", "JobResult.warnings notes Stage 7 (real Maps enrichment) isn't in the job runner's scope -- no fabricated maps_matched/maps_rating fields", "High", "tests/unit/test_job_translation.py"),
    T("JOB-11", "Async Job API", "features.tech_stack=true is a nudge, not a guarantee",
      "A job request with features.tech_stack=true", "Submit such a job",
      "N/A", "Query phrasing nudges the planner toward needs_tech_stack=true, but the result documents this isn't guaranteed", "Medium", "tests/unit/test_job_translation.py"),
    T("JOB-12", "Async Job API", "Multiple target.locations produces a coverage warning",
      "A job request with 2+ entries in target.locations", "Submit such a job",
      "N/A", "Warning that locations are joined into one 'or' phrase, not independently searched per-region", "Medium", "tests/unit/test_job_translation.py"),
    T("JOB-13", "Async Job API", "Clarification-triggering query maps to awaiting_input",
      "A job whose translated query is vague enough to trigger the clarification gate", "Submit such a job",
      "N/A", "JobResult.status='awaiting_input' -- matches the guide's own frontend workflow-state name for this situation", "Medium", "tests/mocked/test_job_lifecycle.py"),
    T("JOB-14", "Async Job API", "Shortfall maps to partially_completed",
      "A job requesting more leads than genuinely exist for the query", "Submit a job with an unrealistic max_prospects for a niche target",
      "N/A", "JobResult.status='partially_completed', counts reflect the real shortfall honestly", "Medium", "tests/mocked/test_job_lifecycle.py::test_shortfall_maps_to_partially_completed"),
    T("JOB-15", "Async Job API", "Moderation-blocked job maps to failed with a validation-class error",
      "A job whose translated query is policy-violating", "Submit such a job",
      "N/A", "JobResult.status='failed', errors[] has an ErrorEnvelope with error_class='validation'", "High", "tests/mocked/test_job_lifecycle.py"),
    T("JOB-16", "Async Job API", "Unexpected exception never crashes the process",
      "Simulate run_pipeline() raising an unhandled exception", "Trigger such a failure (test-only scenario)",
      "N/A", "Job resolves to status='failed' with an ErrorEnvelope (error_class='internal'); the API process itself keeps running normally for other jobs", "Critical", "tests/mocked/test_job_lifecycle.py (unexpected-exception handling test)"),
    T("JOB-17", "Async Job API", "Heartbeat advances independently of progress events",
      "A job whose pipeline stage is slow/stuck on a provider call", "Observe last_heartbeat_at over a 30+ second window with no progress change",
      "N/A", "last_heartbeat_at still advances roughly every 15s even with no stage/progress change -- proves the job is alive, not stalled", "Medium", "None (manual, needs a real slow call) -- design documented in JOB_EXECUTION.md"),
    T("JOB-18", "Async Job API", "Score.state is not_attempted when scoring never ran",
      "A qualified prospect whose relevance_scoring never attached a score (e.g. cached-reuse path)", "Inspect JobResult.prospects[].score for such a case",
      "N/A", "score.value=null, score.state='not_attempted'", "Medium", "tests/mocked/test_job_lifecycle.py::test_score_state_not_attempted_when_classification_missing"),
    T("JOB-19", "Async Job API", "Score.state is unknown for a genuine data-consistency surprise",
      "A qualified prospect with no matching row in summary['results']", "Inspect such a case",
      "N/A", "score.value=null, score.state='unknown'", "Low", "tests/mocked/test_job_lifecycle.py::test_score_state_unknown_when_no_matching_results_row"),
    T("JOB-20", "Async Job API", "evidence_refs use the documented artifact:// scheme",
      "A completed job with qualified prospects", "Inspect JobResult.prospects[].evidence_refs",
      "N/A", "Each ref looks like artifact://storage/<domain>/ -- consistent with docs/backend-handoff/ARTIFACTS_AND_STORAGE.md", "Low", "tests/mocked/test_job_lifecycle.py"),
    T("JOB-21", "Async Job API", "Resume creates a new job reusing the original's target",
      "An original job that has reached a terminal state (failed/cancelled/partially_completed)", "POST /api/v1/jobs/{original_id}/resume with a new job_id + correlation_id",
      "N/A", "202 Accepted; new job's target/limits/features exactly match the original's; new JobResult.resumed_from = original job_id", "High", "tests/mocked/test_api_jobs.py::test_resume_endpoint_reuses_original_target"),
    T("JOB-22", "Async Job API", "Resume skips businesses the original already committed",
      "An original job that committed some businesses before stopping", "Resume it, then check storage/ for re-scraping of those same domains",
      "N/A", "Already-committed domains are NOT re-scraped by the resumed job -- confirms coarse-grained resume via the pipeline's own cache", "High", "None (manual, live) -- design in JOB_EXECUTION.md"),
    T("JOB-23", "Async Job API", "Resume with an unknown original job_id fails",
      "N/A", "POST /api/v1/jobs/does-not-exist/resume", "N/A", "404 Not Found", "High", "tests/mocked/test_api_jobs.py::test_resume_unknown_original_returns_404"),
    T("JOB-24", "Async Job API", "Resume with a new job_id already in use fails",
      "Two existing jobs, A (terminal) and B (any state)", "POST /api/v1/jobs/A/resume with body job_id=B (B's id, already taken)",
      "N/A", "409 Conflict", "Medium", "tests/mocked/test_job_lifecycle.py::test_resume_duplicate_new_job_id_raises"),
    T("JOB-25", "Async Job API", "All 4 job schemas validate real example payloads",
      "N/A", "Validate examples against docs/backend-handoff/schemas/{request,progress,result,error}.schema.json",
      "N/A", "All examples validate cleanly; schema files match the current Pydantic models (no drift)", "Critical", "tests/contract/test_job_schemas.py"),
]

# ===========================================================================
# DATA & STORAGE CONTRACTS
# ===========================================================================

DATA_ROWS = [
    T("DATA-01", "Data & Storage", "crawl_index.csv gets one row per crawled page",
      "A completed run with a multi-page business", "Inspect crawl_index.csv after a run",
      "N/A", "One row per page (not per business); all documented columns populated (company_name, domain, page_url, crawl_status, etc.)", "Critical", "None (manual) -- schema in docs/DATA_CONTRACTS.md"),
    T("DATA-02", "Data & Storage", "txt_path resolves correctly regardless of separator style",
      "Running on Windows (backslash paths) vs. a path with forward slashes", "Load a txt_path value and resolve its basename using the documented [\\\\/] split logic",
      "N/A", "Basename resolves correctly either way -- a real, previously-shipped bug this now guards against", "High", "tests/unit/test_windows_paths.py"),
    T("DATA-03", "Data & Storage", "leads_clean.csv is a correct per-business rollup",
      "A completed run", "Inspect leads_clean.csv",
      "N/A", "One row per business (not per page); email/phone are the union of every value found across ALL that business's pages, deduplicated", "High", "tests/unit/test_page_rollup.py::test_contact_union_across_pages"),
    T("DATA-04", "Data & Storage", "page_title/description picked by name-match, not shortest URL",
      "A business that's a subsection of a larger shared site (e.g. a marina inside a city park authority's site)", "Inspect leads_clean.csv for such a business",
      "N/A", "page_title/description come from whichever page's title best matches the business's own name -- NOT simply the shortest URL", "High", "tests/unit/test_page_rollup.py::test_shared_institutional_domain_prefers_name_matching_page"),
    T("DATA-05", "Data & Storage", "Off-domain email dropped when an own-domain email exists",
      "A business with both an own-domain and unrelated off-domain email found across its pages", "Inspect the email field for such a business",
      "N/A", "The own-domain (or known free-webmail) address is preferred; a genuinely unrelated off-domain address is dropped", "Medium", "tests/unit/test_page_rollup.py::test_off_domain_email_dropped_when_own_domain_email_exists"),
    T("DATA-06", "Data & Storage", "Review-platform JSON always has platform + reviews keys",
      "Any storage/<domain>/reviews/<platform>.json file, including a miss", "Inspect a review file where nothing was found",
      "N/A", "Both 'platform' and 'reviews' keys present even when reviews=[] -- this is what lets accuracy_check.py distinguish real platform records from bookkeeping files", "High", "None (manual) -- schema in docs/DATA_CONTRACTS.md"),
    T("DATA-07", "Data & Storage", "google_maps.json and category.json are excluded from review iteration",
      "A domain with both review platform files and these two bookkeeping files", "Run accuracy_check.py against it",
      "N/A", "Neither file is treated as a review-platform record (correctly lacks the platform+reviews key pair)", "Medium", "None (manual)"),
    T("DATA-08", "Data & Storage", "Domain-key normalization is consistent everywhere",
      "A business referenced sometimes by bare domain, sometimes by full URL", "Cross-check storage/, crawl_index.csv, and the RAG index for the same business",
      "N/A", "All three use the same bare registered domain as the key (via domain_utils.domain_key()) -- never a mix of bare domain and full URL", "Critical", "None (manual) -- a real, previously-fixed bug"),
    T("DATA-09", "Data & Storage", "R2 sync preserves folder hierarchy with correct key format",
      "R2 credentials configured", "Sync a domain's storage/ folder to R2, inspect the bucket structure",
      "N/A", "Remote object keys use forward slashes regardless of host OS (Path.as_posix()) -- a real, previously-fixed Windows bug", "High", "None (manual, needs real R2 bucket) -- design in docs/DATA_CONTRACTS.md"),
    T("DATA-10", "Data & Storage", "R2 sync skips unchanged files on second run",
      "A domain already synced once", "Run the sync again with no changes to that domain's files",
      "N/A", "Second sync is a no-op for unchanged files (incremental, manifest-based skip)", "Medium", "None (manual, needs real R2 bucket)"),
    T("DATA-11", "Data & Storage", "R2 sync is silently disabled without credentials",
      "R2 credentials NOT configured", "Run a normal pipeline query",
      "N/A", "No error, no crash -- sync is simply skipped; pipeline behaves identically to having R2 configured, minus the remote copy", "Medium", "None (manual)"),
]

# ===========================================================================
# ERROR HANDLING & RESILIENCE
# ===========================================================================

ERROR_ROWS = [
    T("ERR-01", "Error Handling & Resilience", "Transient network blip is retried, not treated as permanent",
      "Simulate a brief network interruption during any provider call", "Interrupt network briefly, then let the pipeline continue/retry",
      "N/A", "The affected call is retried automatically rather than caching a false negative", "Critical", "tests/mocked/test_cache_correctness.py (multiple scenarios)"),
    T("ERR-02", "Error Handling & Resilience", "Groq TPM rate limit triggers automatic fallback",
      "Simulate hitting Groq's tokens-per-minute limit", "Trigger a rate-limited call",
      "N/A", "Log shows 'Rate limit reached...'; the secondary model (qwen/qwen3.6-27b) is used automatically -- visible in logs, not a failure requiring action", "High", "None (manual, hard to force in a test env)"),
    T("ERR-03", "Error Handling & Resilience", "ScrapingBee quota exhaustion degrades gracefully",
      "A ScrapingBee plan at its 1000-call monthly limit", "Attempt a premium fetch after quota exhaustion",
      "N/A", "\"Monthly API calls limit reached: 1000\" logged; native-only fetching continues pipeline-wide; not treated as a code bug", "Medium", "None (manual, needs quota actually exhausted)"),
    T("ERR-04", "Error Handling & Resilience", "Every classify_error() input maps to exactly one of the 8 taxonomy classes",
      "N/A", "Feed classify_error() representative failure messages for each class (validation, authentication, quota, rate_limit, transient_provider, blocked_source, data_quality, internal)",
      "Sample messages per class", "Each input maps to the correct, single ErrorClass -- no message falls through to an unexpected default", "High", "tests/unit/test_error_taxonomy.py (if present) / job_contracts.classify_error()"),
    T("ERR-05", "Error Handling & Resilience", "ErrorEnvelope always carries the request's correlation_id",
      "A job that fails for any reason", "Inspect JobResult.errors[0].correlation_id",
      "N/A", "Matches the correlation_id from the original JobRequest exactly -- enables tracing across systems", "High", "tests/mocked/test_job_lifecycle.py"),
]

# ===========================================================================
# SECURITY
# ===========================================================================

SECURITY_ROWS = [
    T("SEC-01", "Security", "Unconfigured server fails closed on every protected route",
      "AIBDM_API_KEY unset on the server", "Call every protected endpoint (/pipeline/run, /leads, /leads/count, /jobs*) with any header",
      "N/A", "All return 503 -- never run open just because no key was configured", "Critical", "None (manual)"),
    T("SEC-02", "Security", "Wrong API key is rejected uniformly",
      "AIBDM_API_KEY configured", "Call protected endpoints with an incorrect key",
      "N/A", "401 on every one; comparison is constant-time (secrets.compare_digest), not vulnerable to a timing side-channel", "High", "None (manual) -- code review of require_api_key()"),
    T("SEC-03", "Security", "No secret ever appears in log output",
      "A full pipeline run with all provider keys configured", "Grep the full log output for the actual configured API key values",
      "N/A", "Zero matches -- no key/token is ever printed, even at DEBUG-adjacent log lines", "Critical", "None (manual, grep-based audit)"),
    T("SEC-04", "Security", "No real .env file is ever committed or archived",
      "N/A", "1. `git ls-files | grep -i \"^\\.env$\"`\n2. Build a handoff ZIP via `git archive` and inspect its contents",
      "N/A", "No bare .env file tracked in git or present in the archive -- only .env.example", "Critical", "None (manual) -- verified repeatedly during this project's handoff work"),
    T("SEC-05", "Security", "A git-archive-built ZIP contains no real credentials",
      "N/A", "`git archive --format=zip -o test.zip HEAD`, then grep the extracted contents for any real-looking secret pattern",
      "N/A", "Only placeholder values (e.g. 'replace_with_local_secret') appear, never a real key", "Critical", "None (manual) -- verified during this project's own handoff packaging"),
    T("SEC-06", "Security", "Scraped page text cannot alter LLM tool behavior (UNTESTED SURFACE)",
      "A business whose page text contains prompt-injection-style content (e.g. \"ignore previous instructions and...\")", "Crawl such a page, then run a find_and_filter query against that business",
      "Adversarial page content", "EXPECTED test to design and run: confirm the LLM's answer/tool-use is NOT hijacked by page content. This is a REAL, currently untested gap -- flagged, not assumed safe. Treat any deviation as a genuine finding, not noise.", "Critical", "None -- this is the one deliberately-flagged gap in this project's own test coverage"),
]

# ===========================================================================
# KNOWN LIMITATIONS (confirm expected behavior, don't file as bugs)
# ===========================================================================

LIMIT_ROWS = [
    T("LIM-01", "Known Limitations", "Yelp/Trustpilot scraping is blocked by bot protection",
      "A business with a real Yelp/Trustpilot listing", "Attempt review harvesting from Yelp or Trustpilot for such a business",
      "N/A", "EXPECTED: bot-protection defeats even the premium fetch tier. Confirm the failure is graceful (no crash), not that it magically starts working.", "Low", "See docs/KNOWN_LIMITATIONS.md"),
    T("LIM-02", "Known Limitations", "Dockwa reviews are never extracted",
      "A marina business with a Dockwa review widget", "Attempt review harvesting from Dockwa",
      "N/A", "EXPECTED: review text is JS-rendered client-side; the scraper only sees the pre-render HTML shell, so extraction finds nothing", "Low", "See docs/KNOWN_LIMITATIONS.md"),
    T("LIM-03", "Known Limitations", "'CRM' doesn't fuzzy-match 'customer relationship management'",
      "A business page that says \"customer relationship management system\" but never the literal word CRM", "Ask a find_and_filter question using the literal term \"CRM\"",
      "N/A", "EXPECTED (deliberate design): no fuzzy synonym matching exists; the query's own keyword-expansion step is expected to spell out realistic surface forms instead", "Low", "See docs/KNOWN_LIMITATIONS.md"),
    T("LIM-04", "Known Limitations", "api.py exposes only Stages 1-2, never Stages 3-8",
      "N/A", "Attempt to reach Stage 4/7/8 functionality via any /api/v1/pipeline/run or /api/v1/jobs/* response",
      "N/A", "EXPECTED: final_reasoning, tech_stack, maps enrichment, accuracy audit results are never present in these API responses -- only main.py's CLI reaches them", "Medium", "See docs/INTEGRATION_NOTES.md, docs/backend-handoff/JOB_EXECUTION.md"),
    T("LIM-05", "Known Limitations", "No artifact is ever content-hashed",
      "Any completed job", "Inspect JobResult.artifact_manifest[].sha256",
      "N/A", "EXPECTED: always null. No file in this codebase is hashed anywhere.", "Low", "See docs/backend-handoff/ARTIFACTS_AND_STORAGE.md"),
    T("LIM-06", "Known Limitations", "Job store and idempotency cache are lost on process restart",
      "A running API server with active jobs", "Restart the uvicorn process, then query a job_id that existed before the restart",
      "N/A", "EXPECTED: 404 -- there is no durable backing store (Redis/DB); everything is an in-process dict", "Medium", "See docs/backend-handoff/JOB_EXECUTION.md"),
    T("LIM-07", "Known Limitations", "Resume restarts from the beginning, not mid-scrape",
      "A resumed job", "Compare a resumed job's discovery-round behavior to a fresh job with the same target",
      "N/A", "EXPECTED: resume re-runs the full discovery loop from round 1 -- it just happens to skip already-committed businesses via the existing cache, it does NOT restore in-flight round state", "Low", "See docs/backend-handoff/JOB_EXECUTION.md"),
    T("LIM-08", "Known Limitations", "robots.txt Disallow rules are never enforced",
      "A site with a Disallow rule on an otherwise-crawlable page", "Crawl that page",
      "N/A", "EXPECTED: the page is fetched anyway. robots.txt is only parsed for Sitemap: discovery. See STG2-13 for the paired functional test.", "Low", "See docs/backend-handoff/KNOWN_LIMITATIONS.md"),
    T("LIM-09", "Known Limitations", "No non-LLM query-planning path exists",
      "N/A", "Attempt to run the pipeline with no Groq key configured at all",
      "N/A", "EXPECTED: pipeline cannot plan or execute ANY query without Groq -- there is no fallback rule-based parser. This is a genuine, permanent architectural fact, not a bug to fix.", "Low", "See docs/backend-handoff/KNOWN_LIMITATIONS.md (the LLM-status discrepancy note)"),
]


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    build_readme(wb)
    write_sheet(wb, "Setup & Environment", "Installing, configuring, and proving a clean environment works.", SETUP_ROWS)
    write_sheet(wb, "Pipeline Stages 1-8", "The core CLI pipeline (main.py), stage by stage. See docs/ARCHITECTURE.md and docs/backend-handoff/PIPELINE_STAGES.md for deeper context.", STAGE_ROWS)
    write_sheet(wb, "Sync API", "The blocking REST endpoints under /api/v1: health, pipeline/run, leads (cursor-paginated).", SYNC_API_ROWS)
    write_sheet(wb, "Async Job API", "The queued job endpoints under /api/v1/jobs: submit, status, result, cancel, resume, Idempotency-Key.", JOB_API_ROWS)
    write_sheet(wb, "Data & Storage Contracts", "What gets written to disk (or R2) and its guarantees. See docs/DATA_CONTRACTS.md.", DATA_ROWS)
    write_sheet(wb, "Error Handling & Resilience", "Provider failures, retries, circuit breakers, the error taxonomy.", ERROR_ROWS)
    write_sheet(wb, "Security", "Auth, secrets, fail-closed behavior -- including one deliberately-flagged untested surface.", SECURITY_ROWS)
    write_sheet(wb, "Known Limitations", "Documented, accepted gaps. Confirm the documented behavior occurs -- don't file these as bugs.", LIMIT_ROWS)

    wb.move_sheet("Read Me", offset=-len(wb.sheetnames))
    out_path = "qa/AI-BDM_QA_Test_Plan.xlsx"
    wb.save(out_path)
    total = sum(len(rows) for rows in [
        SETUP_ROWS, STAGE_ROWS, SYNC_API_ROWS, JOB_API_ROWS, DATA_ROWS, ERROR_ROWS, SECURITY_ROWS, LIMIT_ROWS,
    ])
    print(f"Wrote {out_path} -- {total} test cases across {len(wb.sheetnames) - 1} test tabs.")


if __name__ == "__main__":
    main()
